"""Read-only source inventory; never print personal records."""
from pathlib import Path
from collections import Counter
from zipfile import ZipFile
import json
import xml.etree.ElementTree as ET
import importlib.metadata
import openpyxl

root = Path(__file__).resolve().parents[1]
path = next(root.glob('*.xlsx'))
wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
for ws in wb:
    rows = list(ws.iter_rows(values_only=True))
    headers = rows[0]
    source = [r for r in rows[1:] if any(x is not None and not str(x).startswith('=') for x in r)]
    out = {'sheet': ws.title, 'shape': [ws.max_row, ws.max_column], 'headers': headers, 'non_formula_rows': len(source)}
    if ws.title in ['Data_Penduduk','Data_Pelayanan','Data_UMKM','Potensi_Desa']:
        out['types'] = {str(h): dict(Counter(type(r[i]).__name__ for r in source)) for i,h in enumerate(headers) if h}
        for col in ['dusun','status','jenis_kelamin','jenis_usaha','kategori','kondisi','tanggal','tanggal_lahir','latitude','longitude']:
            if col in headers:
                vals = [r[headers.index(col)] for r in source if r[headers.index(col)] is not None]
                if col in ['tanggal','tanggal_lahir','latitude','longitude']:
                    out[col] = {'min':str(min(vals)), 'max':str(max(vals))} if vals else {}
                else:
                    out[col] = dict(Counter(str(v) for v in vals))
        out['formula_counts'] = {str(h):sum(isinstance(r[i],str) and r[i].startswith('=') for r in rows[1:]) for i,h in enumerate(headers) if h}
    print(json.dumps(out, ensure_ascii=True, default=str))
wb.close()
slide = next(root.glob('*.pptx'))
with ZipFile(slide) as z:
    slides = sorted([n for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml')])
    print('SLIDE_COUNT', len(slides))
    for name in slides[:8]:
        text = [el.text for el in ET.fromstring(z.read(name)).iter() if el.tag.endswith('}t') and el.text]
        print(name, ' | '.join(text)[:900])
print('VERSIONS', {m:importlib.metadata.version(m) for m in ['streamlit','pandas','plotly','openpyxl','pytest','pillow']})
