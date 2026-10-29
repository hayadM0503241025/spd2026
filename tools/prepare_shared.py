"""Fasilitator only: prepare a NEW immutable baseline from the source workbook.

Run: python tools/prepare_shared.py --workbook 'path.xlsx' --age-reference 2026-10-29
Never overwrites a published version. Original workbook and identifiers are not exported.
"""
from pathlib import Path
import argparse
import hashlib
import json
import sys
from datetime import date
import openpyxl
import pandas as pd

sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from studio.constants import DATASETS
from studio.validation import clean_text, ages_on, age_groups, identity_tokens, parse_number, valid_coordinates


def prepare_sheet(rows, dataset, reference):
    headers=[str(c).strip() for c in rows[0]]
    # Some worksheets contain formatting/cells beyond the named source table.
    width=len(headers)
    data=[list(row[:width])+[None]*max(0,width-len(row)) for row in rows[1:]]
    df=pd.DataFrame(data,columns=headers)
    source_columns=[c for c in df if c not in ('usia','kelompok_usia','bulan')]
    # Formula-only placeholders are not records. IDs alone are also insufficient.
    substantive=[c for c in source_columns if c!='ID']
    has_source=df[substantive].map(lambda x: pd.notna(x) and not (isinstance(x,str) and (not x.strip() or x.startswith('=')))).any(axis=1)
    removed=int((~has_source).sum())
    df=df.loc[has_source].reset_index(drop=True)
    reports=[{'level':'info','message':f'{len(df)} record sumber; {removed} baris kosong/rumus saja dikeluarkan.'}]
    for column in df:
        if column not in ('usia','kelompok_usia','bulan'):
            formulas=df[column].map(lambda x:isinstance(x,str) and x.startswith('='))
            if formulas.any():
                reports.append({'level':'warning','message':f'{int(formulas.sum())} rumus pada kolom sumber {column} dikosongkan; cache Excel tidak digunakan.'})
                df.loc[formulas,column]=pd.NA
    if dataset=='penduduk':
        for column,prefix in [('nik','NIK'),('no_kk','KK')]:
            df[column],invalid=identity_tokens(df[column],prefix)
            reports.append({'level':'warning' if invalid else 'info','message':f'{invalid} {column} bukan teks 16 digit: ditandai tidak valid dan dikeluarkan dari hitungan identitas unik. Identitas valid diganti token konsisten.'})
        births=pd.to_datetime(df.tanggal_lahir,errors='coerce')
        df['usia']=ages_on(births,reference)
        df['kelompok_usia']=age_groups(df.usia)
        reports.append({'level':'info','message':f'Usia dan kelompok usia dihitung ulang dari tanggal lahir pada {reference.isoformat()}; formula/cache usia diabaikan. Tanggal lahir dan nama dihapus dari berkas latihan.'})
        df=df.drop(columns=['nama','tanggal_lahir'])
    if dataset=='pelayanan':
        df['tanggal']=pd.to_datetime(df.tanggal,errors='coerce')
        df['bulan']=df.tanggal.dt.to_period('M').astype('string')
        reports.append({'level':'info','message':'Bulan diturunkan ulang dari tanggal sebagai YYYY-MM; tahun yang berbeda tidak digabung.'})
    numbers={'penduduk':['usia'],'pelayanan':['lama_proses'],'umkm':['omset','latitude','longitude'],'potensi':['latitude','longitude']}[dataset]
    for column in numbers:
        before=df[column].notna()
        df[column]=df[column].map(parse_number)
        failed=int((before & df[column].isna()).sum())
        if failed:
            reports.append({'level':'warning','message':f'{failed} nilai gagal dikonversi menjadi angka di {column}.'})
    if {'latitude','longitude'}<=set(df):
        invalid=~valid_coordinates(df)
        reports.append({'level':'warning' if invalid.any() else 'info','message':f'{int(invalid.sum())} pasangan koordinat di luar rentang/tidak lengkap dikosongkan. Tidak dilakukan pembagian atau koreksi dugaan.'})
        df.loc[invalid,['latitude','longitude']]=float('nan')
    for column in df:
        if column not in numbers and column!='tanggal':
            df[column]=df[column].map(clean_text).astype('string')
    duplicate=int(df.duplicated().sum())
    reports.append({'level':'warning' if duplicate else 'info','message':f'{duplicate} dugaan duplikasi record lengkap; tidak dihapus otomatis.'})
    return df,reports,removed


def main():
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--workbook',required=True,type=Path)
    parser.add_argument('--age-reference',required=True,type=date.fromisoformat)
    args=parser.parse_args()
    root=Path(__file__).resolve().parents[1]
    source_hash=hashlib.sha256(args.workbook.read_bytes()).hexdigest()
    version=f'desa-{args.age_reference.isoformat()}-{source_hash[:10]}-p1'
    folder=root/'data'/'shared'/version
    if folder.exists():
        raise SystemExit('Versi sudah tersedia; tidak ditimpa. Gunakan workbook atau revisi aturan persiapan yang baru.')
    workbook=openpyxl.load_workbook(args.workbook,read_only=True,data_only=False)
    manifest={'version':version,'source_file':args.workbook.name,'source_sha256':source_hash,
              'label':'Data latihan bersama dari workbook referensi',
              'processing':{'revision':'p1','age_reference':args.age_reference.isoformat(),
                            'ages':'tahun penuh; 0–14 / 15–59 / 60+',
                            'identities':'teks 16 digit valid saja; token konsisten; numerik tidak valid',
                            'coordinates':'rentang sah saja; tidak ditebak',
                            'formulas':'diabaikan; usia/kelompok/bulan dihitung ulang',
                            'privacy':'nama dan tanggal lahir dihapus; NIK/KK diganti token',
                            'categories':'rapikan spasi; tidak menebak padanan kategori'},'datasets':{}}
    prepared={}
    for dataset,sheet in DATASETS.items():
        try:
            rows=list(workbook[sheet].iter_rows(values_only=True))
            frame,reports,removed=prepare_sheet(rows,dataset,args.age_reference)
            dates=[c for c in frame if pd.api.types.is_datetime64_any_dtype(frame[c])]
            payload=frame.to_csv(index=False,date_format='%Y-%m-%d').encode('utf-8')
            prepared[dataset]=payload
            manifest['datasets'][dataset]={'sheet':sheet,'file':f'{dataset}.csv','rows':len(frame),
                'sha256':hashlib.sha256(payload).hexdigest(), 'dates':dates,
                'dtypes':{c:'Float64' if pd.api.types.is_numeric_dtype(frame[c]) else 'string' for c in frame if c not in dates},
                'reports':reports,'removed_rows':removed,'missing':{c:int(frame[c].isna().sum()) for c in frame},
                'column_mapping':{c:c for c in frame}}
        except (KeyError,ValueError,TypeError) as exc:
            manifest.setdefault('errors',{})[dataset]=f'Gagal menyiapkan sheet: {type(exc).__name__}'
    workbook.close()
    if not prepared:
        raise SystemExit('Tidak ada dataset yang berhasil disiapkan.')
    folder.mkdir(parents=True)
    for dataset,payload in prepared.items():
        (folder/f'{dataset}.csv').write_bytes(payload)
    (folder/'manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
    registry_path=root/'data'/'registry.json'
    registry=json.loads(registry_path.read_text(encoding='utf-8')) if registry_path.exists() else {'versions':{}}
    registry['current']=version
    registry['versions'][version]=version
    registry_path.write_text(json.dumps(registry,indent=2),encoding='utf-8')
    print(json.dumps({'version':version,'rows':{k:v['rows'] for k,v in manifest['datasets'].items()},'errors':manifest.get('errors',{})}))


if __name__=='__main__':
    main()
